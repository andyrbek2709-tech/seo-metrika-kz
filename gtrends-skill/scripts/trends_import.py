#!/usr/bin/env python3
"""
trends_import.py — импорт CSV-экспортов Google Trends → сводный .xlsx.

Google Trends бесплатен, без аккаунта и без карты, но даёт не абсолютный объём,
а ОТНОСИТЕЛЬНЫЙ интерес (0–100) внутри одного сравнения (≤5 фраз). Чтобы сравнивать
фразы из разных батчей, в каждый батч включают одну общую фразу-«анкор»; этот скрипт
пересчитывает (нормирует) каждый файл по анкору, приводя всё к единой шкале.

Принцип WAT: решения (что сравнивать, выводы) — в чате; скрипт только парсит CSV
из trends.google.com (кнопка «Скачать» на графике «Динамика популярности») и пишет .xlsx.

Использование:
  # один батч (фразы сопоставимы внутри файла без нормировки)
  python trends_import.py --csv multiTimeline.csv --out trends.xlsx

  # несколько батчей с общим анкором → кросс-файловая шкала
  python trends_import.py --csv A.csv B.csv C.csv \
      --anchor "проектная документация" --out trends_kz.xlsx
"""

import argparse
import csv
import io
import os
import re
import sys
from typing import Any, Dict, List, Optional


def _norm(s: str) -> str:
    return (s or "").replace("﻿", "").replace("\xa0", " ").strip().lower()


def _read_csv_text(path: str) -> str:
    for enc in ("utf-8-sig", "utf-8", "utf-16", "cp1251"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                return fh.read()
        except (UnicodeError, LookupError):
            continue
    sys.exit(f"[ОШИБКА] Не смог прочитать {path} (кодировка).")


_DATE_HDR = ("неделя", "week", "день", "day", "месяц", "month",
             "время", "time", "дата", "date", "час", "hour")


def parse_trends(path: str) -> Dict[str, float]:
    """Из multiTimeline.csv → {фраза: средний интерес}."""
    rows = list(csv.reader(io.StringIO(_read_csv_text(path))))
    hdr: Optional[int] = None
    # заголовок таблицы — строка с «фраза: (Гео)»
    for i, row in enumerate(rows):
        if any(re.search(r":\s*\(", c) for c in row):
            hdr = i
            break
    if hdr is None:  # запасной вариант: строка, начинающаяся с «Неделя/Week/…»
        for i, row in enumerate(rows):
            if len(row) > 1 and _norm(row[0]) in _DATE_HDR:
                hdr = i
                break
    if hdr is None:
        sys.exit(f"[ОШИБКА] {path}: не нашёл строку-заголовок Trends "
                 f"(ожидал «фраза: (Гео)» или «Неделя,…»). Это multiTimeline.csv?")

    header = rows[hdr]
    terms = [re.sub(r":\s*\(.*\)\s*$", "", c).strip() for c in header[1:]]
    sums = [0.0] * len(terms)
    cnts = [0] * len(terms)
    for row in rows[hdr + 1:]:
        if len(row) < 2:
            continue
        for j in range(len(terms)):
            ci = j + 1
            if ci >= len(row):
                continue
            v = row[ci].strip()
            if not v:
                continue
            if v.replace(" ", "") in ("<1", "<1%"):
                val = 0.5
            else:
                try:
                    val = float(v.replace(",", "."))
                except ValueError:
                    continue
            sums[j] += val
            cnts[j] += 1
    return {terms[j]: (sums[j] / cnts[j] if cnts[j] else 0.0)
            for j in range(len(terms)) if terms[j]}


def import_trends(paths: List[str], anchor: Optional[str]) -> List[Dict[str, Any]]:
    per_term: Dict[str, List[float]] = {}
    for p in paths:
        means = parse_trends(p)
        scale = 1.0
        if anchor:
            akey = next((t for t in means if anchor.lower() in t.lower()), None)
            if akey is None:
                print(f"[ПРЕДУПРЕЖДЕНИЕ] {os.path.basename(p)}: анкор «{anchor}» "
                      f"не найден — файл не нормирован.")
            elif means[akey] <= 0:
                print(f"[ПРЕДУПРЕЖДЕНИЕ] {os.path.basename(p)}: анкор «{akey}» "
                      f"имеет нулевой интерес — файл не нормирован.")
            else:
                scale = 100.0 / means[akey]
        for t, m in means.items():
            per_term.setdefault(t, []).append(m * scale)

    rows = [{
        "запрос": t,
        "относит_интерес": round(sum(v) / len(v), 1),
        "нормировано": "да" if anchor else "нет",
    } for t, v in per_term.items()]
    rows.sort(key=lambda r: -r["относит_интерес"])
    return rows


def write_xlsx(path: str, rows: List[Dict[str, Any]]) -> None:
    if os.path.splitext(path)[1].lower() != ".xlsx":
        sys.exit("[ОШИБКА] --out должен быть .xlsx")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Google Trends KZ"
    if not rows:
        ws["A1"] = "нет данных (проверь, что это multiTimeline.csv)"
        wb.save(path)
        print(f"[OK] Пустой результат → {path}")
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    fill = PatternFill("solid", fgColor="4285F4")
    font = Font(bold=True, color="FFFFFF")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r.get(h) for h in headers])
    for c, h in enumerate(headers, start=1):
        width = max(len(str(h)), *(len(str(r.get(h, ""))) for r in rows))
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 12), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    wb.save(path)
    print(f"[OK] Записано {len(rows)} строк → {path}")


def main():
    ap = argparse.ArgumentParser(description="Импорт Google Trends CSV → сводный .xlsx")
    ap.add_argument("--csv", nargs="+", required=True,
                    help="один или несколько multiTimeline.csv из trends.google.com")
    ap.add_argument("--anchor", default=None,
                    help="фраза-анкор, общая для всех батчей (кросс-файловая нормировка)")
    ap.add_argument("--out", required=True, help="путь к .xlsx")
    args = ap.parse_args()
    write_xlsx(args.out, import_trends(args.csv, args.anchor))


if __name__ == "__main__":
    main()
